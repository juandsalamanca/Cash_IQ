from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pandas as pd
from src.projections import build_projections_table


def collapse_other(df, keep_index, other_name, index_names):
    keep = df.loc[keep_index].copy() if len(keep_index) else df.iloc[0:0].copy()
    other = df.drop(index=keep_index, errors="ignore").copy()
    if len(other):
        other_row = other.sum(axis=0)
        other_idx = pd.MultiIndex.from_tuples([(other_name, "Other", "")], names=index_names)
        other_df = pd.DataFrame([other_row.values], index=other_idx, columns=df.columns)
        keep = pd.concat([keep, other_df], axis=0)
    return keep

def get_cash_balance(total_inflows, total_outflows, beginning_cash_balance, all_week_starts):
    # =========================
    # BEGIN/END CASH BALANCE
    # =========================
    beg_bal_series = pd.Series(index=all_week_starts, dtype=float)
    end_bal_series = pd.Series(index=all_week_starts, dtype=float)

    running_begin = beginning_cash_balance
    for w in all_week_starts:
        beg_bal_series[w] = running_begin
        running_end = running_begin + float(total_inflows[w]) - float(total_outflows[w])
        end_bal_series[w] = running_end
        running_begin = running_end

    return beg_bal_series, end_bal_series

def build_inflows_outflows(combined_full, actual_week_starts, all_week_starts, TOP_N_INFLOW_LINES, TOP_N_OUTFLOW_LINES, idx_names):
    # =========================
    # BUILD INFLOWS/OUTFLOWS PRESENTATION (NO FLAT)
    # =========================
    trailing_actual = combined_full[actual_week_starts].copy()
    inflow_mask  = combined_full.sum(axis=1) > 0
    outflow_mask = combined_full.sum(axis=1) < 0

    # rank lines by trailing magnitude
    top_inflows = (
        combined_full.loc[inflow_mask]
        .assign(trailing_inflow=lambda df: trailing_actual.loc[df.index].clip(lower=0).sum(axis=1))
        .sort_values("trailing_inflow", ascending=False)
        .head(TOP_N_INFLOW_LINES)
        .index
    )

    top_outflows = (
        combined_full.loc[outflow_mask]
        .assign(trailing_outflow=lambda df: (-trailing_actual.loc[df.index].clip(upper=0)).sum(axis=1))
        .sort_values("trailing_outflow", ascending=False)
        .head(TOP_N_OUTFLOW_LINES)
        .index
    )

    in_rows = []
    out_rows = []
    cols = combined_full.columns
    ref = [0 for n in range(len(combined_full.columns))]
    for i in range(len(combined_full)):
        in_row = []
        out_row = []
        idx  = combined_full.index[i]
        for j in range(len(combined_full.columns)):
            value = combined_full.iloc[i, j]
            if value > 0:
                in_row.append(value)
                out_row.append(0)
            elif value < 0:
                in_row.append(0)
                out_row.append(value)
            else:
                in_row.append(0)
                out_row.append(0)

        if in_row != ref:
            in_rows.append(pd.Series(in_row, index=cols, name=idx))
        if out_row != ref:
            out_rows.append(pd.Series(out_row, index=cols, name=idx))
        
    inflows_tbl = pd.DataFrame(in_rows, columns=combined_full.columns)
    inflows_tbl.index.names = ['split_account', 'split_type', 'split_detail_type']
    outflows_tbl = pd.DataFrame(out_rows, columns=combined_full.columns)
    outflows_tbl.index.names = ['split_account', 'split_type', 'split_detail_type']

    #inflows_tbl  = collapse_other(inflows_tbl,  top_inflows,  "Other Inflows",  idx_names)
    #outflows_tbl = collapse_other(outflows_tbl, top_outflows, "Other Outflows", idx_names)

    # Presentation: inflows positive; outflows positive
    inflows_present  = inflows_tbl.copy()
    outflows_present = outflows_tbl.copy().abs()

    total_inflows  = inflows_present.sum(axis=0)
    total_outflows = outflows_present.sum(axis=0)

    print("Inflows present:")
    print(inflows_present.index)
    print("Outflows present:")
    print(outflows_present.index)
    print("-"*100)
    # Account for any empty split account, marked as unmapped
    inflows_present.index = inflows_present.index.set_levels(
        ['Other Inflows' if level == '' else level for level in inflows_present.index.levels[0]],
        level=0
    )
    outflows_present.index = outflows_present.index.set_levels(
        ['Other Outflows' if level == '' else level for level in outflows_present.index.levels[0]],
        level=0
    )
    print("Inflows present:")
    print(inflows_present.index)
    print("Outflows present:")
    print(outflows_present.index)
    print("-"*100)
    return inflows_present, outflows_present, total_inflows, total_outflows

def write_output_excel(all_week_starts, inflows_by_cat, outflows_by_cat, inflows_present, outflows_present, total_inflows,
                        total_outflows, cc_spend_proj_display=None, cc_spend_actual_display=None, cc_payment_alloc_present=None, cc_spend_txn=None, 
                        cc_payment_schedule=None, cc_txn_df_dict=None, beg_bal_series=None, end_bal_series=None, PROJ_WEEK1_START="", OUTPUT_XLSX="", week1_cash_balance=0.0,
                        VENDOR_SUMMARY_PATH=None, ar=None, ar_assumptions_df=None):
    # =========================
    # WRITE OUTPUT EXCEL
    # =========================
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        # Summary
        summary = pd.DataFrame(
            {
                "Week Start": all_week_starts,
                "Beginning Bank Balance": [beg_bal_series[w] for w in all_week_starts],
                "Total Cash Inflows": [total_inflows[w] for w in all_week_starts],
                "Total Cash Outflows": [total_outflows[w] for w in all_week_starts],
                "Ending Bank Balance": [end_bal_series[w] for w in all_week_starts],
            }
        )
        summary.to_excel(writer, sheet_name="Summary", index=False)

        # Cash details
        inflows_present.reset_index().to_excel(writer, sheet_name="Cash Inflows (Detail)", index=False)
        outflows_present.reset_index().to_excel(writer, sheet_name="Cash Outflows (Detail)", index=False)

        # Credit card sheets
        cc_list = [cc_spend_txn, cc_spend_actual_display, cc_spend_proj_display, cc_payment_schedule, cc_payment_alloc_present]
        if all(cc is not None for cc in cc_list):
            cc_spend_txn.sort_values(["account_name","date"]).to_excel(writer, sheet_name="CC Spend - Transactions", index=False)
            cc_spend_actual_display.reset_index().to_excel(writer, sheet_name="CC Spend - Weekly (Hist)", index=False)
            cc_spend_proj_display.reset_index().to_excel(writer, sheet_name="CC Spend - Weekly (Proj)", index=False)
            cc_payment_schedule.to_excel(writer, sheet_name="CC Payments - Schedule", index=False)
            cc_payment_alloc_present.reset_index().to_excel(writer, sheet_name="Cash - CC Pay Allocation", index=False)

        if cc_txn_df_dict is not None:
            for cc_df_name in cc_txn_df_dict:
                cc_df = cc_txn_df_dict[cc_df_name]
                cc_df.to_excel(writer, sheet_name=cc_df_name, index=False)
        
        proj_sheet, inflow_section_indexes, outflow_section_indexes, cash_balance_indexes = build_projections_table(all_week_starts, 
                                                                                                                    inflows_by_cat, 
                                                                                                                    outflows_by_cat, 
                                                                                                                    beg_bal_series, 
                                                                                                                    end_bal_series, 
                                                                                                                    total_inflows, 
                                                                                                                    total_outflows, 
                                                                                                                    inflows_present, 
                                                                                                                    outflows_present, 
                                                                                                                    week1_cash_balance)

        proj_sheet.to_excel(writer, sheet_name="Projections (Table)", index=False)

        ar_list = [ar, ar_assumptions_df]
        if all(a is not None for a in ar_list):
            ar.to_excel(writer, sheet_name="AR Aging (Raw)", index=False)
            ar_assumptions_df.to_excel(writer, sheet_name="AR Collections (Assumptions)", index=False)

        if VENDOR_SUMMARY_PATH is not None:
            vendor_summary = pd.read_excel(VENDOR_SUMMARY_PATH)
            vendor_summary.columns = [str(c).strip() for c in vendor_summary.columns]
            vendor_summary.to_excel(writer, sheet_name="Expenses by Vendor (Raw)", index=False)

        print(f"Saved: {OUTPUT_XLSX}")
        print(f"Projection Week 1 starts: {PROJ_WEEK1_START.date()} (Monday)")

    return inflow_section_indexes, outflow_section_indexes, cash_balance_indexes

def calculate_category_totals(OUTPUT_XLSX, inflows_by_cat, outflows_by_cat):

    _, cash_balance_indexes, inflow_section_indexes, outflow_section_indexes = get_category_indexes(OUTPUT_XLSX, 
                                                                                                              inflows_by_cat, 
                                                                                                              outflows_by_cat, 
                                                                                                              sheet_name="Projections (Table)")
    wb = load_workbook(OUTPUT_XLSX, data_only=True)
    wb_2 = load_workbook(OUTPUT_XLSX, data_only=True)
    ws = wb["Projections (Table)"]
    ws_2 = wb_2["Projections (Table)"]
    # -----------------------------------------
    # Get the category totals
    # -----------------------------------------

    for section_indexes in [inflow_section_indexes, outflow_section_indexes]:

        if section_indexes is inflow_section_indexes:
            operation = 'SUM'
        else:
            operation = '-SUM'

        for i in range(len(section_indexes)):
            idx = section_indexes[i]
            if idx == section_indexes[-1]:
                break

            next_idx = section_indexes[i+1]

            row = ws[idx]
            for col in range(3, len(row)):
                col_letter = get_column_letter(col+1)
                if next_idx-2 >= idx+1:
                    row[col].value = f'={operation}({col_letter}{idx+1}:{col_letter}{next_idx-2})'
                    number_value = 0
                    for n in range(idx+1, next_idx-1):
                        summand = ws[n][col].value
                        if summand is not None and pd.isna(summand) == False:
                            try:
                                if operation == 'SUM':
                                    number_value += summand
                                else:
                                    number_value -= summand
                            except Exception as e:
                                raise TypeError(f"Error adding {summand} from row {n} to total for category at row {idx}, column {col}: {str(e)}")
                    ws_2[idx][col].value = number_value
                else:
                    row[col].value = 0.0
                    ws_2[idx][col].value = 0.0


    # -----------------------------------------
    # Calculate total inflows and outflows
    # -----------------------------------------

    total_inflows_row_idx = inflow_section_indexes[-1]
    total_outflows_row_idx = outflow_section_indexes[-1]
    for col in range(3, ws.max_column):
        #print(col)
        col_letter = get_column_letter(col+1)
        # Total Inflows
        inflows_sum_String =  f'='
        number_value = 0
        for i in range(len(inflow_section_indexes)-1):
            inflows_sum_String += f'{col_letter}{inflow_section_indexes[i]}+'
            sumand = ws_2[inflow_section_indexes[i]][col].value
            if sumand is not None and pd.isna(sumand) == False:
                number_value += sumand
        inflows_sum_String = inflows_sum_String.rstrip('+')
        ws_2[total_inflows_row_idx][col].value = number_value
        row = ws[total_inflows_row_idx]
        row[col].value = inflows_sum_String

        # Total Outflows
        outflows_sum_String =  f'='
        number_value = 0
        for i in range(len(outflow_section_indexes)-1):
            outflows_sum_String += f'{col_letter}{outflow_section_indexes[i]}+'
            summand = ws_2[outflow_section_indexes[i]][col].value
            if summand is not None and pd.isna(summand) == False:
                number_value += summand
        outflows_sum_String = outflows_sum_String.rstrip('+')
        ws_2[total_outflows_row_idx][col].value = number_value
        row = ws[total_outflows_row_idx]
        row[col].value = outflows_sum_String


    # ----------------------------------------------
    # Calculate beginning and ending cash balances
    # ----------------------------------------------

    beg_cash_row_idx = cash_balance_indexes[0]
    end_cash_row_idx = cash_balance_indexes[1]
    beg_row = ws[beg_cash_row_idx]
    end_row = ws[end_cash_row_idx]
    beg_row_2 = ws_2[beg_cash_row_idx]
    end_row_2 = ws_2[end_cash_row_idx]

    # Logic is different for the values of the past than teh ones of the present:

    # We need to have the hard coded numbers done first, that way we don't run into strings when doing the calculations


    # Hard coded numbers

    # Cash balances for the past
    for col in range(6, 2, -1):
        col_letter = get_column_letter(col+1)
        next_col_letter = get_column_letter(col+2)

        # End balance is just the beg balanace from teh next column
        end_row_2[col].value = ws_2[beg_cash_row_idx][col+1].value

        # Beg balaance is end balaance - inflows - outflows (already negative)
        beg_row_2[col].value = ws_2[end_cash_row_idx][col].value - ws_2[total_outflows_row_idx][col].value - ws_2[total_inflows_row_idx][col].value

    # Cash balances for the future
    for col in range(7, ws.max_column):
        col_letter = get_column_letter(col+1)
        prev_col_letter = get_column_letter(col)

        # Beg balaance is end balaance from previous column except for the one of the present
        if col == 7:
            pass
        else:
            beg_row_2[col].value = ws_2[end_cash_row_idx][col-1].value

        # End balance is beg balance + inflows + outflows (already negative)
        end_row_2[col].value = ws_2[beg_cash_row_idx][col].value + ws_2[total_inflows_row_idx][col].value + ws_2[total_outflows_row_idx][col].value


    # Excel formulas

    # Cash balances for the past
    for col in range(3, 7):
        col_letter = get_column_letter(col+1)
        next_col_letter = get_column_letter(col+2)

        # End balance is just the beg balanace from teh next column
        end_row[col].value = f'={next_col_letter}{beg_cash_row_idx}'

        # Beg balaance is end balaance - inflows - outflows (already negative)
        beg_row[col].value = f'={col_letter}{end_cash_row_idx}-{col_letter}{total_outflows_row_idx}-{col_letter}{total_inflows_row_idx}'

    # Cash balances for the future
    for col in range(7, ws.max_column):
        col_letter = get_column_letter(col+1)
        prev_col_letter = get_column_letter(col)

        # Beg balaance is end balaance from previous column except for the one of the present
        if col == 7:
            pass
        else:
            beg_row[col].value = f'={prev_col_letter}{end_cash_row_idx}'

        # End balance is beg balance + inflows + outflows (already negative)
        end_row[col].value = f'={col_letter}{beg_cash_row_idx}+{col_letter}{total_inflows_row_idx}+{col_letter}{total_outflows_row_idx}'


    wb.save(OUTPUT_XLSX)
    TEMP_OUTPUT_XLSX = str(OUTPUT_XLSX)[:-5] + "_temp.xlsx"
    wb_2.save(TEMP_OUTPUT_XLSX)
    wb.close()
    wb_2.close()
    return TEMP_OUTPUT_XLSX


def get_category_indexes(excel_path, inflows_by_cat, outflows_by_cat, sheet_name="", final_format=False):

    wb = load_workbook(excel_path)

    if sheet_name:
        ws = wb[sheet_name]
    else:
        ws = wb.active
    cat_indexes = {}

    if final_format:
        cat_range = range(1,3)
    else:
        cat_range = range(0,2)

    for i in range(1, ws.max_row+1):
        for j in cat_range:
            if ws[i][j].value is not None:
                cat_indexes[ws[i][j].value] = [i,j]

    cash_balance_indexes = []
    inflow_section_indexes = []
    outflow_section_indexes = []
    for key in cat_indexes:
        if key in ["Beginning Bank Balance", "Ending Bank Balance"]:
            cash_balance_indexes.append(cat_indexes[key][0])
        if key == "Total Cash Inflows" or key in inflows_by_cat:
            inflow_section_indexes.append(cat_indexes[key][0])
        if key == "Total Cash Outflows" or key in outflows_by_cat:
            outflow_section_indexes.append(cat_indexes[key][0])

    wb.close()

    return cat_indexes, cash_balance_indexes, inflow_section_indexes, outflow_section_indexes