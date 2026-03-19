from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Border, Side
from src.general_postprocessing import get_category_indexes

def style_projections(OUTPUT_XLSX, inflows_by_cat, outflows_by_cat, cash_floor):

    _, cash_balance_indexes, inflow_section_indexes, outflow_section_indexes = get_category_indexes(OUTPUT_XLSX, 
                                                                                                              inflows_by_cat, 
                                                                                                              outflows_by_cat, 
                                                                                                              sheet_name="Projections (Table)")

    header_rows = 2

    wb = load_workbook(OUTPUT_XLSX)

    # Define styles
    sheet_name = "Projections (Table)"

    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found")

    ws = wb[sheet_name]

    # Add the cash floor row
    for _ in range(header_rows):
        ws.insert_rows(1)

    ws[1][0].value = "13 Week Cash Flow Forcast"
    ws[2][2].value = "Cash Floor"
    ws[2][3].value = cash_floor

    for idx_list in [inflow_section_indexes, outflow_section_indexes, cash_balance_indexes]:
        for i in range(len(idx_list)):
            idx_list[i] = idx_list[i] + header_rows

    # Styles
    

    title_fill = PatternFill(
        start_color="F2977E",
        end_color="F2977E",
        fill_type="solid"
    )

    for cell in ws[1]:
        cell.fill = title_fill

    header_fill = PatternFill(
        start_color="A3A5D0",
        end_color="A3A5D0",
        fill_type="solid"
    )

    accounting_format = '_($* #,##0.00_);_($* (#,##0.00);_($* "-"??_);_(@_)'

    font_style = Font(name="Aptos Narrow", size=12)
    bold_font_style = Font(name="Aptos Narrow", size=12, bold=True)

    # Apply styles
    for i, row in enumerate(ws.iter_rows()):
        for j, cell in enumerate(row):
            # Bold for columns A and B and row 4
            # Because here we use i index from the enumerate method use 0 based index instead of 1 based like we do when using ws[i]
            # So for the fourth row we use i==3 instead of i==4
            if j==0 or j==1 or i==3:
                cell.font = bold_font_style
            else:
                cell.font = font_style

            # Accounting format only for numbers
            if isinstance(cell.value, (int, float)):
                cell.number_format = accounting_format

    # Header fill
    for i in range(3, 5):
        for cell in ws[i]:
            cell.fill = header_fill

    # Category fill
    category_fill = PatternFill(
        start_color="BFBFBF",
        end_color="BFBFBF",
        fill_type="solid"
    )

    # Projection fill
    projection_fill = PatternFill(
        start_color="53C9B8",
        end_color="53C9B8",
        fill_type="solid"
    )

    for section_indexes in [inflow_section_indexes, outflow_section_indexes]:

        for i in range(len(section_indexes)):
            idx = section_indexes[i]
            row = ws[idx]

            # Total inflows and outflows are one column to the left of the categories
            if i == len(section_indexes)-1:
                start = 0
            else:
                start = 1
            # Color categories with grey
            for j in range(start, len(row)):
                row[j].fill = category_fill

            # Color categories with tiel
            if i != len(section_indexes)-1:
                for idx in range(section_indexes[i]+1, section_indexes[i+1]):
                    row = ws[idx]
                    for j in range(7, len(row)):
                        row[j].fill = projection_fill

    # Apply color to bag end cash
    beg_cash_row_idx = cash_balance_indexes[0]
    end_cash_row_idx = cash_balance_indexes[1]
    beg_row = ws[beg_cash_row_idx]
    end_row = ws[end_cash_row_idx]

    for col in range(ws.max_column):
        beg_row[col].fill = header_fill
        end_row[col].fill = header_fill

    # Add conditional formatting to end balance
    conditional_font_color = Font(color="9C0006")
    conditional_fill = PatternFill(
        start_color="FFC7CE",
        end_color="FFC7CE",
        fill_type="solid"
    )
    rule = FormulaRule(
        formula=[f"H{end_cash_row_idx}<D{header_rows}"],
        font=conditional_font_color,
        fill=conditional_fill
        )

    ws.conditional_formatting.add(f"H{end_cash_row_idx}:T{end_cash_row_idx}", rule)

    # Fix row 3: change text values and remove borders
    no_border = Side(style=None)
    for j, cell in enumerate(ws[3]):
        if j == 0:
            cell.value = "Date at Start of Week"
        if j in [1,2]:
            cell.value = ""
        cell.border = Border(
            left=no_border,
            right=no_border,
            top=no_border,
            bottom=no_border)

    # Add hard border between past weeks and projections
    med = Side(style="medium")      # or "medium", "thick"
    #border = Border(right=thin)

    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
        cell = row[6]  # Column G (0-based index → A=0, G=6)
        
        # Preserve existing borders
        cell.border = Border(
            left=cell.border.left,
            right=med,
            top=cell.border.top,
            bottom=cell.border.bottom
        )

    # Hide all tabs except for projections
    for sheet in wb.worksheets:
        if sheet.title != sheet_name:
            sheet.sheet_state = "hidden"
        else:
            sheet.sheet_state = "visible"

    # Freze panes at D5
    ws.freeze_panes = "D5"

    wb.save(OUTPUT_XLSX)
    wb.close()